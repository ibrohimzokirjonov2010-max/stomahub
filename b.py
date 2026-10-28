from openpyxl import Workbook
from openpyxl.styles import PatternFill,Font
from openpyxl.utils import get_column_letter as L
wb=Workbook();Y=PatternFill("solid",fgColor="FFFF00");B=Font(bold=True)
f=wb.active;f.title="Farazlar"
f.append(["StomaHub moliyaviy modeli — FARAZLAR. DIQQAT: barcha qiymatlar TAXMINIY (namuna) ma'lumot, haqiqiy emas! Sariq kataklarni haqiqiy ma'lumot bilan to'ldiring."]);f["A1"].font=B
f.append(["Ko'rsatkich","Qiymat","Birlik","Belgi"])
A=[("gmv","Marketplace: 1-oy savdo hajmi (GMV)",50000000,"UZS/oy"),
("gmvg","Marketplace: GMV oylik o'sishi",0.15,"%"),
("kom","Marketplace komissiyasi",0.10,"%"),
("sem","Seminarlar soni oyiga",2,"dona"),
("tick","Bir seminarga chiptalar",30,"dona"),
("tp","Chipta narxi",500000,"UZS"),
("sh","Platforma ulushi chiptadan",0.30,"%"),
("cl0","CRM: boshlang'ich klinikalar",10,"dona"),
("cln","CRM: oyiga yangi klinikalar",8,"dona"),
("ch","CRM: oylik churn",0.03,"%"),
("crm","CRM obuna narxi",300000,"UZS/oy"),
("app0","Ilova: 1-oy premium foydalanuvchilar",100,"dona"),
("appg","Ilova: foydalanuvchilar oylik o'sishi",0.20,"%"),
("appp","Ilova premium obuna narxi",30000,"UZS/oy"),
("team","Jamoa (ish haqi) xarajati",40000000,"UZS/oy"),
("ads","Reklama xarajati",15000000,"UZS/oy"),
("host","Hosting/serverlar",3000000,"UZS/oy"),
("other","Boshqa doimiy xarajatlar",2000000,"UZS/oy"),
("pay","To'lov tizimi komissiyasi (Payme/Click/Uzum)",0.015,"% daromaddan"),
("tax","Aylanmadan olinadigan soliq stavkasi",0.04,"% daromaddan — joriy stavkani tekshiring! IT-Park rezidentligi imtiyozli rejim berishi mumkin"),
("cash0","Boshlang'ich kapital (investitsiya)",200000000,"UZS")]
R={}
for i,(k,n,v,u) in enumerate(A,3):
  f.append([n,v,u,"TAXMINIY"]);f.cell(i,2).fill=Y;R[k]=f"Farazlar!$B${i}"
  if "%" in u:f.cell(i,2).number_format="0.0%"
  else:f.cell(i,2).number_format="#,##0"
f.column_dimensions["A"].width=48;f.column_dimensions["B"].width=16;f.column_dimensions["C"].width=60
def sheet(t,rows):
  s=wb.create_sheet(t);s.append([t+" (TAXMINIY namuna ma'lumotlar asosida)"]);s["A1"].font=B
  s.append(["Ko'rsatkich"]+[f"{m}-oy" for m in range(1,13)]+["Jami"])
  for r,(n,fn) in enumerate(rows,3):
    s.cell(r,1,n)
    for m in range(1,13):
      c=L(m+1);p=L(m);s.cell(r,m+1,"="+fn(c,p,m,r)).number_format="#,##0"
    s.cell(r,14,f"=SUM(B{r}:M{r})").number_format="#,##0"
  s.column_dimensions["A"].width=34;return s
g=R
sheet("Daromad",[
("Marketplace GMV",lambda c,p,m,r:g['gmv'] if m==1 else f"{p}{r}*(1+{g['gmvg']})"),
("Marketplace komissiya daromadi",lambda c,p,m,r:f"{c}3*{g['kom']}"),
("Ta'lim/seminar daromadi",lambda c,p,m,r:f"{g['sem']}*{g['tick']}*{g['tp']}*{g['sh']}"),
("CRM faol klinikalar",lambda c,p,m,r:f"{g['cl0']}+{g['cln']}" if m==1 else f"{p}{r}*(1-{g['ch']})+{g['cln']}"),
("CRM obuna daromadi",lambda c,p,m,r:f"{c}6*{g['crm']}"),
("Ilova premium foydalanuvchilar",lambda c,p,m,r:g['app0'] if m==1 else f"{p}{r}*(1+{g['appg']})"),
("Mobil ilova daromadi",lambda c,p,m,r:f"{c}8*{g['appp']}"),
("JAMI DAROMAD",lambda c,p,m,r:f"{c}4+{c}5+{c}7+{c}9")])
sheet("Xarajatlar",[
("Jamoa",lambda *a:g['team']),("Reklama",lambda *a:g['ads']),("Hosting/serverlar",lambda *a:g['host']),
("Boshqa doimiy",lambda *a:g['other']),
("To'lov komissiyalari (Payme/Click/Uzum)",lambda c,p,m,r:f"Daromad!{c}10*{g['pay']}"),
("Aylanma soliq (stavkani tekshiring; IT-Park varianti)",lambda c,p,m,r:f"Daromad!{c}10*{g['tax']}"),
("JAMI XARAJAT",lambda c,p,m,r:f"SUM({c}3:{c}8)")])
sheet("P&L oylik",[
("Daromad",lambda c,p,m,r:f"Daromad!{c}10"),
("O'zgaruvchan xarajatlar (komissiya+soliq)",lambda c,p,m,r:f"Xarajatlar!{c}7+Xarajatlar!{c}8"),
("Yalpi foyda",lambda c,p,m,r:f"{c}3-{c}4"),
("Doimiy xarajatlar",lambda c,p,m,r:f"SUM(Xarajatlar!{c}3:{c}6)"),
("SOF FOYDA",lambda c,p,m,r:f"{c}5-{c}6")])
s=sheet("Pul oqimi",[
("Boshlang'ich qoldiq",lambda c,p,m,r:g['cash0'] if m==1 else f"{p}7"),
("Kirim (daromad)",lambda c,p,m,r:f"Daromad!{c}10"),
("Chiqim (xarajatlar)",lambda c,p,m,r:f"Xarajatlar!{c}9"),
("Sof pul oqimi",lambda c,p,m,r:f"{c}4-{c}5"),
("Yakuniy qoldiq",lambda c,p,m,r:f"{c}3+{c}6"),
("Kumulyativ sof foyda",lambda c,p,m,r:f"'P&L oylik'!{c}7" if m==1 else f"{p}8+'P&L oylik'!{c}7")])
for cc in ("N3","N7","N8"):s[cc]=None
for cc in ("N3","N6","N8"):wb["Daromad"][cc]=None
fb=wb.create_sheet("Foyda nuqtasi");fb.column_dimensions["A"].width=70;fb.column_dimensions["B"].width=20
P="'P&L oylik'"
rows=[("FOYDA NUQTASI (TAXMINIY namuna ma'lumotlar asosida)",None),
("Birinchi foydali oy (oylik sof foyda ≥ 0)",f'=IFERROR(MATCH(TRUE,INDEX({P}!B7:M7>=0,0),0)&"-oy","12 oy ichida yo\'q")'),
("Investitsiya qoplanadigan oy (kumulyativ foyda ≥ 0)",f'=IFERROR(MATCH(TRUE,INDEX(\'Pul oqimi\'!B8:M8>=0,0),0)&"-oy","12 oy ichida yo\'q")'),
("Eng past pul qoldig'i (12 oy)","=MIN('Pul oqimi'!B7:M7)"),
("Oylik doimiy xarajatlar",f"={g['team']}+{g['ads']}+{g['host']}+{g['other']}"),
("O'zgaruvchan xarajat ulushi (to'lov+soliq)",f"={g['pay']}+{g['tax']}"),
("Kerakli oylik daromad (foyda nuqtasi)","=B5/(1-B6)"),
("Faqat CRM bilan qoplash uchun kerakli klinikalar",f"=ROUNDUP(B5/({g['crm']}*(1-B6)),0)"),
("Faqat marketplace bilan qoplash uchun kerakli oylik GMV",f"=B5/({g['kom']}*(1-B6))"),
("Faqat seminar bilan: kerakli chiptalar soni oyiga",f"=ROUNDUP(B5/({g['tp']}*{g['sh']}*(1-B6)),0)"),
("12-oy daromadi (taqqoslash uchun)","=Daromad!M10")]
for i,(a,b) in enumerate(rows,1):
  fb.cell(i,1,a);
  if b: fb.cell(i,2,b).number_format="#,##0"
fb["B6"].number_format="0.0%";fb["A1"].font=B
wb.save("StomaHub_moliyaviy_model.xlsx")
